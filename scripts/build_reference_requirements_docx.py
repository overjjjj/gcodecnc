from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "参考系统功能需求确认清单.md"
OUTPUT = ROOT / "docs" / "参考系统功能需求确认清单.docx"
TITLE = "加工中心控制系统功能需求确认清单"
SUBTITLE = "重点：第七章设计界面与 CAM 功能 - 待截图核验版本"

BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
GRID = "B7C4D6"
CONTENT_WIDTH_DXA = 14400  # Letter landscape with 0.5 in margins


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=60, start=90, bottom=60, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, size, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_cell_text(cell, text, header=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or len(text) < 13 else WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, 0, 0, 1.0)
    run = p.add_run(text)
    set_run_font(run, 7.2 if not header else 7.6, bold=header, color="FFFFFF" if header else None)


def get_widths(columns):
    if columns == 6:
        return [900, 2350, 5300, 2400, 1800, 1650]
    if columns == 5:
        return [900, 2650, 6100, 2600, 2150]
    if columns == 4:
        return [1700, 4400, 4700, 3600]
    if columns == 3:
        return [2400, 6900, 5100]
    if columns == 2:
        return [3200, 11200]
    return [CONTENT_WIDTH_DXA // columns] * columns


def split_markdown_row(line):
    return [piece.strip() for piece in line.strip().strip("|").split("|")]


def is_separator(line):
    return bool(re.match(r"^\|(?:\s*:?-{3,}:?\s*\|)+$", line))


def add_table(doc, lines):
    rows = [split_markdown_row(line) for line in lines if not is_separator(line)]
    if not rows:
        return
    columns = len(rows[0])
    table = doc.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    widths = get_widths(columns)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            add_cell_text(cell, value, header=(row_index == 0))
            if row_index == 0:
                set_cell_shading(cell, BLUE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "FAFCFF")
        if row_index == 0:
            set_repeat_table_header(row)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in [
        ("Heading 1", 15, "2E74B5", 14, 7),
        ("Heading 2", 12.5, "2E74B5", 11, 6),
        ("Heading 3", 11, BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(TITLE)
    set_run_font(run, 8, color="6B7A90")
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("需求确认基线 - 参考说明书第七章 - 第 ")
    set_run_font(run, 8, color="6B7A90")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    run = p.add_run(" 页")
    set_run_font(run, 8, color="6B7A90")


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 8, 3, 1.0)
    run = p.add_run(TITLE)
    set_run_font(run, 22, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 0, 12, 1.0)
    run = p.add_run(SUBTITLE)
    set_run_font(run, 10, color="52627A")
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, (label, value) in zip(table.rows, [
        ("基线资料", "《加工中心控制系统用户使用手册 manual_zh_CN》"),
        ("用途", "功能范围、实现优先级、后续软件截图及验收证据的逐项确认"),
        ("版本", "v0.2 - 2026-08-18"),
    ]):
        add_cell_text(row.cells[0], label, header=False)
        add_cell_text(row.cells[1], value, header=False)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            run.font.bold = True
    set_table_geometry(table, [2600, 11800])
    doc.add_paragraph()


def main():
    global SOURCE, OUTPUT, TITLE, SUBTITLE
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--title", default=TITLE)
    parser.add_argument("--subtitle", default=SUBTITLE)
    args = parser.parse_args()
    SOURCE = args.source
    OUTPUT = args.output
    TITLE = args.title
    SUBTITLE = args.subtitle
    source_lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    setup_styles(doc)
    add_header_footer(doc)
    add_title(doc)

    index = 0
    while index < len(source_lines):
        line = source_lines[index].rstrip()
        if not line or line == "---":
            index += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            run = p.add_run(line[2:])
            set_run_font(run, 8.5, color="52627A")
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(source_lines) and source_lines[index].startswith("|"):
                table_lines.append(source_lines[index])
                index += 1
            add_table(doc, table_lines)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                index += 1
                continue
            style = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}.get(level, "Heading 3")
            p = doc.add_paragraph(style=style)
            p.add_run(text)
            index += 1
            continue
        p = doc.add_paragraph()
        set_paragraph_spacing(p, 0, 4, 1.15)
        run = p.add_run(line)
        set_run_font(run, 9.2)
        index += 1

    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = "CNEXT-CAM 项目"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
