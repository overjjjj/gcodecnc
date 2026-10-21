from pathlib import Path
import re
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT=Path(r'D:\yanggao\gcodecnc\整理副本_反编译源码'); UI=ROOT/'ui整理'
classes=['MainWindow','AppGlobal','CatDesign','CatWork','WorkParamForm','DialogEyeletHole','CatWork_Waiting']
counts=Counter(); files={}
for p in (ROOT/'DTM'/'decompiled_app').glob('*.cpp'):
    t=p.read_text(encoding='utf-8',errors='replace')
    for c in classes:
        n=len(re.findall(r'\b'+c+r'::[A-Za-z_]\w*\s*\(',t))
        if n: counts[c]+=n; files.setdefault(c,[]).append(p.name)
doc=Document(); s=doc.sections[0]; s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
doc.styles['Normal'].font.name='Calibri'; doc.styles['Normal'].font.size=Pt(10.5)
for n,z,col in [('Heading 1',16,'2E74B5'),('Heading 2',13,'2E74B5')]:
    st=doc.styles[n]; st.font.name='Calibri'; st.font.size=Pt(z); st.font.color.rgb=RGBColor.from_string(col)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('DTM UI 分层整理与项目组任务下发方案'); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=RGBColor(31,77,120)
p=doc.add_paragraph('开发经理版｜主界面 → 子界面 → 后续 UI → 算法适配'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('一、项目目标与边界',1); doc.add_paragraph('将反编译恢复包整理为项目组可分工、可验证、可重实现的依据。原始目录保持不变，中文注释与分类文件位于 ui整理/；反编译代码不可直接视为原始可编译工程。')
doc.add_heading('二、界面层级与整理顺序',1)
levels=[('第1层：主界面','MainWindow、AppGlobal','应用生命周期、菜单工具栏、页面容器、全局状态和任务入口；代码见 decompiled_0001/0002/0005/0010.cpp。'),('第2层：设计子界面','CatDesign','模型/几何设计、选择集和设计状态；代码见 decompiled_0002/0003/0010.cpp。'),('第2层：加工子界面','CatWork','加工任务编排、队列、取消和进度；代码见 decompiled_0002/0010.cpp。'),('第3层：参数界面','WorkParamForm','刀具、进给、余量、安全高度等参数录入校验；分布于 decompiled_0001 至 0011.cpp。'),('第3层：专用对话框','DialogEyeletHole、CatWork_Waiting','孔特征参数、等待/进度提示、异常确认；代码见 decompiled_0002/0009/0012.cpp。')]
for h,k,b in levels: doc.add_heading(h,2); doc.add_paragraph(k+'：'+b)
doc.add_heading('三、主要函数中文注释与函数规模',1); doc.add_paragraph('整理副本已在文件头和可识别类成员函数前加入中文推断性注释。后续需用测试把推断升级为契约，不能把反编译变量名直接当作业务字段。')
t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,v in zip(t.rows[0].cells,['类/模块','函数数','建议责任']): c.text=v
for c in classes:
    if counts[c]: x=t.add_row().cells; x[0].text=c; x[1].text=str(counts[c]); x[2].text='入口/状态/参数/对话流程复核与单测'
doc.add_heading('四、耦合性审查',1); doc.add_paragraph('函数级清单已统计静态引用频次。高引用函数是候选枢纽，跨 UI 与算法边界的调用按高风险处理。拆分原则：UI 不直接操作几何内存；通过 DTO 和应用服务交互；算法入口统一错误码、取消令牌和进度回调。')
doc.add_heading('五、项目组任务下发',1)
tasks=[('T1 基线固化','平台组','索引、文件映射、样例和构建说明','原始校验值不变且可追溯'),('T2 主界面骨架','UI组长','重建 MainWindow/AppGlobal 生命周期、页面容器和状态机','可启动、可切页、有日志'),('T3 设计子界面','设计组','重建 CatDesign 状态、选择集和刷新','操作可回放、状态有测试'),('T4 加工子界面','加工组','重建 CatWork 任务队列、取消、进度','任务可取消、异常可恢复'),('T5 参数对话框','参数/特征组','重建 WorkParamForm、孔对话框和校验','非法参数阻断、单位明确'),('T6 算法适配','算法组','封装孔、槽框、刀路、仿真入口','黄金样例误差达标'),('T7 集成开源审查','架构/QA','联调、性能、许可证、SBOM、发布清单','CI与回归报告齐全')]
t=doc.add_table(rows=1,cols=4)
for c,v in zip(t.rows[0].cells,['任务','负责人','内容','验收']): c.text=v
for row in tasks:
    x=t.add_row().cells
    for i,v in enumerate(row): x[i].text=v
doc.add_heading('六、后续模块开发方案',1)
for h,b in [('UI层','Qt 页面只展示与转发事件，Presenter/ViewModel 管理状态、忙碌、错误和空状态。'),('应用服务层','编排设计、加工、导出、仿真，提供取消、超时、进度和审计日志。'),('领域模型层','定义模型、特征、刀具、工艺参数、刀路和仿真状态 DTO。'),('算法适配层','为 MyToolPath、ModuleAutoHole、ModuleAutoSlotFrame、GeometryOcct、MillSim 建立版本化接口。'),('基础设施层','封装文件、日志、配置、OpenCascade、G代码解析和线程。')]: doc.add_heading(h,2); doc.add_paragraph(b)
doc.add_heading('七、可复用开源库与选型决策',1)
doc.add_paragraph('以下选型以“优先复用、隔离适配、许可证可审计”为原则；不得直接复制反编译代码或未知来源代码。')
t=doc.add_table(rows=1,cols=4)
for c,v in zip(t.rows[0].cells,['库/项目','拟复用位置','许可证关注点','决策']): c.text=v
libs=[('Qt 5.8 / qtbase','UI、事件循环、线程、网络、SQLite 驱动','按 Qt 版本选择 LGPL/GPL/商业许可；锁定 5.8 ABI','保留现有运行时，先不升级'),('OCCT 7.6','几何、拓扑、STEP/IGES、布尔、可视化','LGPL-2.1 + OCCT exception；保留替换库能力','核心依赖，建立 GeometryAdapter'),('Qt Solutions QtPropertyBrowser','WorkParamForm 属性编辑器','BSD-3-Clause；保留版权和许可证','优先复用'),('QScintilla','G代码/脚本编辑器','GPLv3 或商业许可；需做发布模式决策','隔离为可选插件'),('Botan/OpenSSL/libssh2','密钥、SSH、TLS','逐库保留 LICENSE/NOTICE，避免静态链接不明','先做依赖取证再封装'),('OpenCAMLib','Drop/Push cutter 刀路基线','LGPL；需做算法误差和性能回归','仅作为算法参考/适配器'),('FreeCAD','CAM 数据模型和工作流参考','LGPL-2.1；逐文件核对','参考架构，不直接拷贝'),('LinuxCNC','RS274/G代码语义和测试样例','GPL/LGPL 混合，谨慎复制','只作协议/测试参考')]
for row in libs:
    x=t.add_row().cells
    for i,v in enumerate(row): x[i].text=v
doc.add_heading('八、详细系统设计',1)
for h,b in [('组件边界','DTM UI 进程负责交互和任务编排；dmtube 作为算法服务/动态库边界；通过稳定 DTO 传递几何、特征、刀具、刀路和仿真数据。'),('核心接口','IModelImporter、IFeatureRecognizer、IToolpathGenerator、IGCodePostprocessor、ISimulator、ITaskScheduler；每个接口返回 Result<T>、错误码、进度和取消句柄。'),('数据流','导入模型→几何校验→特征识别→参数校验→策略选择→刀路生成→安全检查→后处理→仿真→报告/导出。'),('状态机','任务状态统一为 Created、Validating、Running、Cancelling、Succeeded、Failed、Cancelled；UI 只订阅状态事件，不直接修改算法状态。'),('线程与性能','几何和刀路在工作线程执行；UI 线程只做轻量更新；对大模型采用增量网格、包围盒预筛和缓存；记录耗时、内存和失败样例。'),('安全设计','所有文件路径、参数范围、G代码词法和刀路越界必须校验；禁止未经确认的 rapid move、主轴/冷却输出进入生产后处理。')]: doc.add_heading(h,2); doc.add_paragraph(b)
doc.add_heading('九、测试与验收设计',1)
for s in ['单元测试：参数边界、单位换算、几何公差、G代码模态状态。','契约测试：UI 与应用服务、应用服务与 dmtube 的 DTO/错误码/取消语义。','算法回归：孔、槽框、轮廓、岛屿、螺旋、清角和三维曲面黄金样例。','安全测试：非法路径、异常模型、自交轮廓、越界刀路、危险 M/G 指令。','性能测试：模型规模、刀路点数、生成耗时、峰值内存和取消响应时间。','发布验收：CI 可重复构建、依赖 SBOM、许可证清单、NOTICE、源码提供义务和可替换动态库验证。']:
    doc.add_paragraph(s,style='List Bullet')
doc.add_heading('十、里程碑、风险与开源要求',1); doc.add_paragraph('M1 基线与依赖清单；M2 主界面可启动；M3 设计/加工页面可操作；M4 参数与算法模拟闭环；M5 刀路/G代码回归；M6 开源发布候选。风险：类型和资源缺失、算法不可验证、全局状态高耦合、许可证不明、Qt/QScintilla/OCCT 组合许可冲突。发布前必须完成权属确认、依赖 SBOM、许可证扫描、NOTICE、敏感信息清理和法务复核。')
doc.add_heading('十一、交付物',1); doc.add_paragraph('分类代码副本、整理清单.csv、函数级清单_含耦合.csv、开源依赖与 GitHub 候选评估.md、本方案 Word、接口契约、测试样例、回归报告、SBOM 和许可证清单。')
out=UI/'DTM_UI分层整理与项目组任务下发方案.docx'; doc.save(out); print(out)
